# -*- coding: utf-8 -*-
"""관광벤처창업론 · 예비관광벤처 사업계획서 주차별 작업 워크북 — 공통 도구

mkworkbook.py 가 `from wb_lib import *` 로 부른다. 둘은 같은 폴더에 있어야 한다.
(2026-09-07 복원 — 원본이 세션 임시 폴더에 있다가 사라져 저장소로 옮겨 왔다.)
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'NanumBarunGothic'
NAVY = '17365D'
HEAD_BG = 'EAF0F7'
ACCENT_BG = 'FDF3E3'
AI_BG = 'F0F4EC'
AI_TAB = 'DCE6D5'
BORDER = 'B7C3D0'
GRAY = '6B7280'
CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT
TOP = WD_ALIGN_VERTICAL.TOP

doc = Document()

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.4)
sec.top_margin = sec.bottom_margin = Cm(1.1)

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.3

W = Cm(18.2)   # 본문 폭


def run(p, text, size=10.5, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(text='', size=10.5, bold=False, color=None, before=0, after=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, color)
    return p


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), color)
    tcPr.append(el)


def _borders(cell, color=BORDER, sz=5):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:color'), color)
        b.append(e)
    tcPr.append(b)


def fill(cell, text='', size=10.0, bold=False, color=None, align=None,
         valign=WD_ALIGN_VERTICAL.CENTER, bg=None, lines=0):
    _borders(cell)
    if bg:
        _shade(cell, bg)
    cell.vertical_alignment = valign
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.25
    if align:
        p.alignment = align
    first = True
    for line in str(text).split('\n'):
        if not first:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            if align:
                p.alignment = align
        first = False
        if line:
            run(p, line, size, bold, color)
    for _ in range(lines):
        q = cell.add_paragraph()
        q.paragraph_format.space_before = Pt(0)
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = 1.7
        run(q, '', size)


# 표가 본문 폭(18.2cm)과 정확히 같으면 워드에서 테두리 두께만큼 오른쪽으로 밀려 나간다.
# 폭을 조금 줄이는 대신 셀 안쪽 여백을 함께 줄여, 글자가 들어갈 자리는 그대로 두었다.
TABLE_MAX = 17.9          # cm — 본문 폭 18.2 에서 0.3 을 여유로 남긴다
CELL_MAR = 57             # twips(0.1cm) — 워드 기본값 108 보다 좁게


def _set_tblpr(tbl, total_cm):
    """tblPr 의 자식은 순서가 정해져 있다 — tblW → jc → tblInd → tblLayout → tblCellMar → tblLook.
    python-docx 가 이미 만들어 둔 tblW 를 고쳐 쓰고, 나머지는 tblLook 앞에 순서대로 끼운다."""
    tblPr = tbl.tblPr

    tw = tblPr.find(qn('w:tblW'))                 # 있는 것을 고친다(새로 붙이면 중복된다)
    if tw is None:
        tw = OxmlElement('w:tblW')
        tblPr.insert(0, tw)
    tw.set(qn('w:type'), 'dxa')
    tw.set(qn('w:w'), str(int(round(total_cm * 567))))

    anchor = tblPr.find(qn('w:tblLook'))
    def put(tag, attrs, children=None):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), str(v))
        for ctag, cattrs in (children or []):
            c = OxmlElement(ctag)
            for k, v in cattrs.items():
                c.set(qn(k), str(v))
            el.append(c)
        if anchor is not None:
            anchor.addprevious(el)
        else:
            tblPr.append(el)

    put('w:tblInd', {'w:type': 'dxa', 'w:w': 0})          # 왼쪽으로 밀리지 않게
    put('w:tblLayout', {'w:type': 'fixed'})
    put('w:tblCellMar', {}, [('w:' + s, {'w:type': 'dxa', 'w:w': v})
                             for s, v in (('left', CELL_MAR), ('right', CELL_MAR),
                                          ('top', 0), ('bottom', 0))])


# 표가 본문 폭(18.2cm)과 정확히 같으면 워드에서 테두리 두께만큼 오른쪽으로 밀려 나간다.
# 폭을 조금 줄이는 대신 셀 안쪽 여백을 함께 줄여, 글자가 들어갈 자리는 그대로 두었다.
TABLE_MAX = 17.9          # cm — 본문 폭 18.2 에서 0.3 을 여유로 남긴다
CELL_MAR = 57             # twips(0.1cm) — 워드 기본값 108 보다 좁게


def table(widths):
    total = sum(w.cm for w in widths)
    k = TABLE_MAX / total if total > TABLE_MAX else 1.0    # 비율은 그대로, 전체만 줄인다

    tw = [int(round(w.cm * k * 567)) for w in widths]      # twips 로 먼저 확정하고
    tw[-1] += int(round(min(total, TABLE_MAX) * 567)) - sum(tw)   # 반올림 오차는 끝 칸이 흡수
    widths = [Cm(v / 567) for v in tw]                     # 그 값에서 Cm 을 만든다

    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _set_tblpr(t._tbl, sum(tw) / 567)

    grid = t._tbl.find(qn('w:tblGrid'))         # tcW 와 tblGrid 를 같은 값으로 맞춘다
    if grid is not None:
        for gc, v in zip(grid.findall(qn('w:gridCol')), tw):
            gc.set(qn('w:w'), str(v))

    t._widths = widths
    t._tw = tw                                  # tcW 를 twips 그대로 쓰려고 남겨 둔다
    return t


def row(t, cells, h=None, exact=False):
    r = t.add_row()
    if h:
        trPr = r._tr.get_or_add_trPr()
        hr = OxmlElement('w:trHeight')
        hr.set(qn('w:val'), str(int(h * 20)))
        hr.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
        trPr.append(hr)
    for i, spec in enumerate(cells):
        c = r.cells[i]
        c.width = t._widths[i]
        tcW = c._tc.get_or_add_tcPr().find(qn('w:tcW'))   # Cm 왕복 반올림으로 1 twip 이 어긋난다
        if tcW is not None:
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(t._tw[i]))
        if isinstance(spec, dict):
            d = dict(spec)
            fill(c, d.pop('text', ''), **d)
        else:
            fill(c, spec)
    return r


def spacer(pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run(p, '', pts)


def page_break():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ── 시트 구성요소 ──────────────────────────────────────────────

def band(left, right):
    t = table([Cm(12.6), Cm(5.6)])
    row(t, [
        {'text': left, 'size': 11.5, 'bold': True, 'color': 'FFFFFF', 'bg': NAVY},
        {'text': right, 'size': 8.5, 'bold': True, 'color': 'FFFFFF', 'bg': NAVY,
         'align': RIGHT},
    ], h=24)
    spacer(4)


def namebar():
    t = table([Cm(6.0), Cm(6.1), Cm(6.1)])
    row(t, [
        {'text': '학번 ________________', 'size': 8.5, 'align': CENTER},
        {'text': '이름 ________________', 'size': 8.5, 'align': CENTER},
        {'text': '작성일 ______ / ______', 'size': 8.5, 'align': CENTER},
    ], h=15)
    spacer(4)


def goal(text):
    t = table([W])
    row(t, [{'text': text, 'size': 9.0, 'color': '33475B', 'bg': 'F7F9FC'}], h=18)
    spacer(4)


def head(text):
    p = para(text, 9.5, True, NAVY, before=2, after=1)
    return p


def star(text, lines=2):
    t = table([W])
    row(t, [{'text': '★  ' + text, 'size': 9.5, 'bold': True, 'color': '9A5B00',
             'bg': ACCENT_BG, 'lines': lines, 'valign': TOP}], h=22)
    spacer(4)


def ai_block(allow, forbid, prompt, check=None):
    t = table([Cm(2.2), Cm(16.0)])
    row(t, [
        {'text': 'AI', 'size': 9.0, 'bold': True, 'color': '2F5D2F', 'bg': AI_TAB,
         'align': CENTER},
        {'text': '써도 되는 것  ·  ' + allow, 'size': 8.5, 'bg': AI_BG},
    ], h=15)
    row(t, [
        {'text': '사용', 'size': 9.0, 'bold': True, 'color': '2F5D2F', 'bg': AI_TAB,
         'align': CENTER},
        {'text': '직접 써야 하는 것  ·  ' + forbid, 'size': 8.5, 'bg': AI_BG,
         'color': 'A03030'},
    ], h=15)
    row(t, [
        {'text': '규칙', 'size': 9.0, 'bold': True, 'color': '2F5D2F', 'bg': AI_TAB,
         'align': CENTER},
        {'text': '프롬프트 예시  ·  ' + prompt, 'size': 8.5, 'bg': AI_BG, 'color': GRAY},
    ], h=15)
    if check:
        row(t, [
            {'text': '검증', 'size': 9.0, 'bold': True, 'color': '2F5D2F', 'bg': AI_TAB,
             'align': CENTER},
            {'text': check, 'size': 8.5, 'bg': AI_BG, 'color': '2F5D2F', 'bold': True},
        ], h=15)
    t2 = table([Cm(2.2), Cm(5.4), Cm(5.3), Cm(5.3)])
    row(t2, [
        {'text': '사용 기록', 'size': 8.5, 'bold': True, 'color': '2F5D2F', 'bg': AI_TAB,
         'align': CENTER},
        {'text': '무엇을 물었나', 'size': 8.5, 'bold': True, 'color': NAVY, 'bg': HEAD_BG,
         'align': CENTER},
        {'text': '받은 답 중 실제로 쓴 것', 'size': 8.5, 'bold': True, 'color': NAVY,
         'bg': HEAD_BG, 'align': CENTER},
        {'text': '내가 고치거나 버린 것', 'size': 8.5, 'bold': True, 'color': NAVY,
         'bg': HEAD_BG, 'align': CENTER},
    ], h=14)
    row(t2, ['', '', '', ''], h=30, exact=True)
    spacer(4)


def submit(text):
    t = table([Cm(2.6), Cm(15.6)])
    row(t, [
        {'text': '다음 주까지', 'size': 8.5, 'bold': True, 'color': 'FFFFFF', 'bg': '4A6C8C',
         'align': CENTER},
        {'text': text, 'size': 9.0, 'bg': 'F7F9FC'},
    ], h=17)
